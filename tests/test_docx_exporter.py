"""Regression coverage for the editable Word export."""

from __future__ import annotations

import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from app import create_app
from core.docx_exporter import DocxExportError, DocxExporter
from pipeline import OUTPUT_FORMATS, PipelineManager
from web.schemas import OutputRequest


def _node(node_id: str, node_type: str, text: str, **attributes: object) -> dict[str, object]:
    return {
        "id": node_id,
        "type": node_type,
        "source": f"source {node_id}",
        "translated_text": text,
        "unit_ids": [f"unit-{node_id}"],
        "attributes": attributes,
    }


def _state() -> dict[str, object]:
    nodes = [
        _node("title", "heading", "一本书", level=1, epub_role="title"),
        _node(
            "copyright",
            "paragraph",
            "版权页第一段\n\n版权页第二段",
            chapter_file="OEBPS/Text/copyright.xhtml",
            source_block_range={"start": 1, "end": 2, "count": 2},
        ),
        _node("toc-heading", "heading", "目录", level=1, epub_role="contents"),
        _node(
            "toc-entries",
            "paragraph",
            "I/75 前言\n\n2 第一章 14",
            structure_type="toc_entry",
            source_block_range={"start": 1, "end": 2, "count": 2},
        ),
        _node("chapter", "heading", "第一章", level=1, epub_role="chapter"),
        _node("body", "paragraph", "正文第一段。"),
        _node("list", "list", "- 第一项\n- 第二项", ordered=False),
        _node("ordered-list", "list", "2. 第二项\n3. 第三项", ordered=True),
        _node("table", "table", "| 名称 | 值 |\n| --- | --- |\n| 项目 | 1972 |"),
        _node(
            "mixed-table",
            "table",
            "说明文字\n| A | B |\n| --- | --- |\n| 1 | 2 |",
        ),
        _node("quote", "blockquote", "> 引用内容\n> 第二行"),
        _node("code", "code", "```text\nprint('hello')\n```"),
        _node("footnote", "paragraph", "注释 I/75，1972 年，第 5—120 页。", structure_type="footnote"),
        _node("after-note", "paragraph", "注释之后的正文。"),
    ]
    units = [
        {
            "id": unit_id,
            "order": order,
            "source": str(node["source"]),
            "translation": str(node["translated_text"]),
            "status": "passed",
        }
        for order, node in enumerate(nodes, 1)
        for unit_id in node["unit_ids"]
    ]
    ids = [str(node["id"]) for node in nodes]
    return {
        "project": {"source_file": {"name": "Example.epub"}},
        "config": {"target_language": "简体中文"},
        "document": {
            "schema_version": 2,
            "format": "epub",
            "source_name": "Example.epub",
            "source_sha256": "source-hash",
            "parts": [
                {"type": "unit", "unit_id": unit["id"]}
                for unit in units
            ],
            "nodes": nodes,
            "reading_order": ids,
            "assets": [],
        },
        "units": units,
        "output": {"format": None, "artifacts": {}},
        "events": [],
    }


class DocxExporterTests(unittest.TestCase):
    def test_workbench_offers_word_and_confirms_the_selected_format(self) -> None:
        root = Path(__file__).resolve().parents[1]
        workbench = (root / "static" / "index.html").read_text(encoding="utf-8")
        script = (root / "static" / "app.js").read_text(encoding="utf-8")
        self.assertIn('<option value="docx">Word (.docx)</option>', workbench)
        self.assertIn('docx: "Word"', script)
        self.assertIn("/static/app.js?v=20260928-docx-r1", workbench)

    def test_exports_editable_a4_structure_with_toc_footnotes_and_trace_map(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            exporter = DocxExporter(temporary)
            metadata = exporter.export(_state())
            target = Path(temporary) / metadata["path"]
            trace_path = Path(temporary) / metadata["trace_map_path"]

            self.assertEqual(target.name, "Example_translated.docx")
            self.assertTrue(trace_path.is_file())
            self.assertEqual(metadata["included_unit_count"], 14)
            self.assertEqual(metadata["size_bytes"], target.stat().st_size)
            with zipfile.ZipFile(target) as archive:
                self.assertIsNone(archive.testzip())

            document = Document(io.BytesIO(target.read_bytes()))
            section = document.sections[0]
            self.assertAlmostEqual(section.page_width.inches, 8.27, delta=0.02)
            self.assertAlmostEqual(section.page_height.inches, 11.69, delta=0.02)
            self.assertAlmostEqual(section.left_margin.inches, 0.75, delta=0.01)
            self.assertEqual(section.start_type, WD_SECTION.NEW_PAGE)

            by_text = {paragraph.text: paragraph for paragraph in document.paragraphs if paragraph.text}
            self.assertEqual(by_text["一本书"].style.name, "Title")
            self.assertIsNone(document.styles["Title"].element.pPr.find(qn("w:pBdr")))
            self.assertEqual(by_text["第一章"].style.name, "Heading 1")
            self.assertTrue(by_text["第一章"].paragraph_format.page_break_before)
            self.assertTrue(by_text["版权页第一段"].paragraph_format.page_break_before)
            self.assertEqual(by_text["版权页第一段"].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertIn("I/75 前言", [p.text for p in document.paragraphs])
            self.assertFalse(by_text["I/75 前言"].paragraph_format.page_break_before)
            self.assertIn("2 第一章 14", [p.text for p in document.paragraphs])
            self.assertEqual(
                [p.text for p in document.paragraphs if p.style.name == "List Bullet"],
                ["第一项", "第二项"],
            )
            self.assertEqual(
                [p.text for p in document.paragraphs if p.style.name == "List Number"],
                ["第二项", "第三项"],
            )
            start_overrides = [
                item.get(qn("w:val"))
                for item in document.part.numbering_part.element.iter(qn("w:startOverride"))
            ]
            self.assertIn("2", start_overrides)
            self.assertIn("3", start_overrides)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(
                [[cell.text for cell in row.cells] for row in document.tables[0].rows],
                [["名称", "值"], ["项目", "1972"]],
            )
            self.assertTrue(any(
                "说明文字" in paragraph.text
                and "| A | B |" in paragraph.text
                and "| 1 | 2 |" in paragraph.text
                for paragraph in document.paragraphs
            ))
            footnote = by_text["注释 I/75，1972 年，第 5—120 页。"]
            self.assertEqual(footnote.style.name, "Footnote Text")
            self.assertEqual(by_text["引用内容"].style.name, "Quote")
            self.assertEqual(by_text["第二行"].style.name, "Quote")
            self.assertEqual(by_text["print('hello')"].style.name, "Code")
            self.assertLess(
                [p.text for p in document.paragraphs].index(footnote.text),
                [p.text for p in document.paragraphs].index("注释之后的正文。"),
            )

            trace = json.loads(trace_path.read_text(encoding="utf-8"))
            self.assertEqual(trace["unit_count"], 14)
            self.assertEqual(
                [item["node_id"] for item in trace["nodes"]],
                [str(node["id"]) for node in _state()["document"]["nodes"]],
            )
            mapped_units = [unit_id for item in trace["nodes"] for unit_id in item["unit_ids"]]
            self.assertEqual(len(mapped_units), len(set(mapped_units)))
            self.assertEqual(set(mapped_units), {unit["id"] for unit in _state()["units"]})

    def test_render_failure_does_not_replace_existing_docx_or_trace_map(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            exporter = DocxExporter(temporary)
            target = exporter._target_path(_state())
            target.parent.mkdir(parents=True)
            target.write_bytes(b"previous docx")
            trace_path = target.with_name(f"{target.name}.map.json")
            trace_path.write_text("previous trace", encoding="utf-8")

            with patch.object(exporter, "_render_payload", side_effect=ValueError("render failed")):
                with self.assertRaises(DocxExportError):
                    exporter.export(_state())

            self.assertEqual(target.read_bytes(), b"previous docx")
            self.assertEqual(trace_path.read_text(encoding="utf-8"), "previous trace")
            self.assertEqual(list(target.parent.glob("*.tmp")), [])

    def test_trace_publish_failure_rolls_back_the_previous_docx(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            exporter = DocxExporter(temporary)
            target = exporter._target_path(_state())
            target.parent.mkdir(parents=True)
            target.write_bytes(b"previous docx")
            trace_path = target.with_name(f"{target.name}.map.json")
            trace_path.write_text("previous trace", encoding="utf-8")
            path_replace = Path.replace

            def fail_trace_publish(source: Path, destination: Path) -> Path:
                if source.name.endswith(".map.json.tmp"):
                    raise PermissionError("simulated trace publish failure")
                return path_replace(source, destination)

            with patch.object(Path, "replace", fail_trace_publish):
                with self.assertRaises(DocxExportError):
                    exporter.export(_state())

            self.assertEqual(target.read_bytes(), b"previous docx")
            self.assertEqual(trace_path.read_text(encoding="utf-8"), "previous trace")
            self.assertEqual(list(target.parent.glob(".*.tmp")), [])

    def test_pipeline_status_and_api_download_support_docx(self) -> None:
        self.assertEqual(OUTPUT_FORMATS, ("markdown", "text", "pdf", "epub", "docx"))
        for output_format in OUTPUT_FORMATS:
            self.assertEqual(OutputRequest(format=output_format).format, output_format)

        with tempfile.TemporaryDirectory() as temporary:
            manager = PipelineManager(Path(temporary) / "runtime")
            manager.state = _state()
            manager._glyph_precheck_locked = lambda: {"available": True}  # type: ignore[method-assign]
            client = TestClient(create_app(manager), base_url="http://127.0.0.1:4873")
            headers = {
                "X-Mode2-Token": client.get("/api/session").json()["token"],
                "Origin": "http://127.0.0.1:4873",
            }

            generated = client.post("/api/project/output", json={"format": "docx"}, headers=headers)
            self.assertEqual(generated.status_code, 200, generated.text)
            self.assertEqual(generated.json()["output"]["filename"], "Example_translated.docx")
            self.assertTrue(generated.json()["readiness"]["formats"]["docx"]["available"])

            downloaded = client.get("/api/project/output/latest", params={"format": "docx"})
            self.assertEqual(downloaded.status_code, 200)
            self.assertEqual(
                downloaded.headers["content-type"],
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            self.assertTrue(downloaded.content.startswith(b"PK"))


if __name__ == "__main__":
    unittest.main()
