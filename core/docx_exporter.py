"""Create an editable Word document from assembled translated nodes."""

from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any

from .assembler import DocumentAssembler
from .epub_layout import build_epub_layout, source_toc_level, split_translated_paragraphs
from .exceptions import PipelineError
from .utils import now_iso

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:  # Keep non-DOCX features usable until the dependency is installed.
    Document = None  # type: ignore[assignment,misc]
    WD_STYLE_TYPE = None  # type: ignore[assignment,misc]
    WD_CELL_VERTICAL_ALIGNMENT = None  # type: ignore[assignment,misc]
    WD_ALIGN_PARAGRAPH = None  # type: ignore[assignment,misc]
    OxmlElement = None  # type: ignore[assignment,misc]
    qn = None  # type: ignore[assignment,misc]
    Inches = Pt = RGBColor = None  # type: ignore[assignment,misc]


class DocxExportError(PipelineError):
    """The project cannot be written as an editable Word document."""


_INLINE_MARKUP = re.compile(
    r"(?P<strong>\*\*.+?\*\*|__.+?__)|(?P<code>`[^`]+`)|(?P<emphasis>\*[^*\n]+\*|_[^_\n]+_)"
)
_LIST_MARKER = re.compile(r"^\s*(?P<marker>(?:[-+*])|(?:\d+[.)]))\s+(?P<text>.*)$")
_TABLE_DIVIDER_CELL = re.compile(r":?-{3,}:?")


class DocxExporter:
    """Write a reflowable A4 DOCX while preserving source order and trace IDs."""

    def __init__(self, runtime_dir: Path | str) -> None:
        self.runtime_dir = Path(runtime_dir).resolve()
        self.output_root = (self.runtime_dir / "output").resolve()
        self.assembler = DocumentAssembler(self.runtime_dir)

    def export(self, state: dict[str, Any]) -> dict[str, Any]:
        try:
            nodes = self.assembler.assemble_nodes(state)
            trace = self.assembler.trace_map(state)
            payload = self._render_payload(nodes, state)
            self._validate_package(payload)

            target = self._target_path(state)
            target.parent.mkdir(parents=True, exist_ok=True)
            trace_path = target.with_name(f"{target.name}.map.json")
            docx_temporary = target.with_name(f".{target.name}.tmp")
            trace_temporary = trace_path.with_name(f".{trace_path.name}.tmp")
            rollback_temporary = target.with_name(f".{target.name}.rollback.tmp")
            had_previous_docx = target.is_file()
            published_docx = False
            try:
                docx_temporary.write_bytes(payload)
                trace_temporary.write_text(
                    json.dumps(trace, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                    newline="\n",
                )
                if had_previous_docx:
                    shutil.copyfile(target, rollback_temporary)
                # Both complete files are staged before either published path is
                # replaced, so a render or staging error leaves the old export intact.
                docx_temporary.replace(target)
                published_docx = True
                trace_temporary.replace(trace_path)
            except OSError:
                if published_docx:
                    if had_previous_docx and rollback_temporary.is_file():
                        rollback_temporary.replace(target)
                    else:
                        target.unlink(missing_ok=True)
                raise
            finally:
                docx_temporary.unlink(missing_ok=True)
                trace_temporary.unlink(missing_ok=True)
                rollback_temporary.unlink(missing_ok=True)

            return {
                "format": "docx",
                "path": target.relative_to(self.runtime_dir).as_posix(),
                "trace_map_path": trace_path.relative_to(self.runtime_dir).as_posix(),
                "filename": target.name,
                "sha256": hashlib.sha256(payload).hexdigest(),
                "exported_at": now_iso(),
                "included_unit_count": len(state.get("units") or []),
                "size_bytes": target.stat().st_size,
            }
        except DocxExportError:
            raise
        except (OSError, ValueError, zipfile.BadZipFile) as exc:
            raise DocxExportError(f"DOCX 生成失败：{exc}") from exc

    def _render_payload(self, nodes: list[dict[str, Any]], state: dict[str, Any]) -> bytes:
        if Document is None:
            raise DocxExportError("Word 导出需要 python-docx，请按 requirements.txt 安装依赖。")

        document = Document()
        section = document.sections[0]
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)

        self._configure_styles(document)
        document.core_properties.title = self.assembler.source_title(state)
        document.core_properties.author = "Mode2 AutoTranslator"

        layout = build_epub_layout(nodes)
        page_breaks = self._page_break_positions(nodes, layout.roles)
        for index, node in enumerate(nodes):
            node_type = str(node.get("type") or "paragraph")
            if node_type == "page_break":
                if index > 0:
                    document.add_page_break()
                continue
            if index in page_breaks and index > 0:
                # The page_break_before property avoids a stray blank paragraph
                # while keeping the target heading editable as a normal style.
                page_break_before = True
            else:
                page_break_before = False
            role = layout.roles.get(str(node.get("id") or ""), "body")
            self._append_node(
                document,
                node,
                role=role,
                page_break_before=page_break_before,
            )

        stream = io.BytesIO()
        document.save(stream)
        return stream.getvalue()

    def _configure_styles(self, document: Any) -> None:
        styles = document.styles
        normal = styles["Normal"]
        self._style_font(normal, size=11, bold=False)
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.35
        normal.paragraph_format.widow_control = True

        self._style_font(styles["Title"], size=22, bold=True)
        styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
        styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        styles["Title"].paragraph_format.space_before = Pt(24)
        styles["Title"].paragraph_format.space_after = Pt(20)
        styles["Title"].paragraph_format.keep_with_next = True
        title_properties = styles["Title"].element.get_or_add_pPr()
        title_border = title_properties.find(qn("w:pBdr"))
        if title_border is not None:
            title_properties.remove(title_border)

        heading_sizes = {1: 16, 2: 14, 3: 12, 4: 11.5, 5: 11, 6: 11}
        for level, size in heading_sizes.items():
            style = styles[f"Heading {level}"]
            self._style_font(style, size=size, bold=True)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_before = Pt(17 if level == 1 else 12)
            style.paragraph_format.space_after = Pt(7)
            style.paragraph_format.keep_with_next = True

        self._add_paragraph_style(document, "Contents Title", size=16, bold=True, before=8, after=12)
        self._add_paragraph_style(document, "Contents Entry", size=10.5, before=0, after=4)
        self._add_paragraph_style(document, "Contents Entry 2", size=10, before=0, after=3, left=0.25)
        self._add_paragraph_style(document, "Front Matter", size=10, before=0, after=7)
        self._add_paragraph_style(document, "Footnote Text", size=9, before=0, after=4, left=0.25)
        self._add_paragraph_style(document, "Quote", size=10.5, before=5, after=8, left=0.3, right=0.2)
        self._add_paragraph_style(document, "Code", size=9, before=2, after=5, left=0.2)

    def _add_paragraph_style(
        self,
        document: Any,
        name: str,
        *,
        size: float,
        bold: bool = False,
        before: float = 0,
        after: float = 0,
        left: float = 0,
        right: float = 0,
    ) -> Any:
        styles = document.styles
        style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        self._style_font(style, size=size, bold=bold)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        if left:
            style.paragraph_format.left_indent = Inches(left)
        if right:
            style.paragraph_format.right_indent = Inches(right)
        style.paragraph_format.widow_control = True
        return style

    @staticmethod
    def _style_font(style: Any, *, size: float, bold: bool) -> None:
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:ascii"), "Aptos")
        rfonts.set(qn("w:hAnsi"), "Aptos")
        rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    def _append_node(
        self,
        document: Any,
        node: dict[str, Any],
        *,
        role: str,
        page_break_before: bool,
    ) -> None:
        metadata = node.get("attributes") if isinstance(node.get("attributes"), dict) else {}
        node_type = str(node.get("type") or "paragraph")
        text = str(node.get("translated_text") or "")
        if not text.strip() and node_type == "image":
            text = str(node.get("source") or "")
        source_classes = self._source_style_classes(metadata)

        if node_type == "heading":
            text = re.sub(r"^\s*#{1,6}\s+", "", text)
            style = self._heading_style(metadata, role, source_classes)
            paragraph = document.add_paragraph(style=style)
            if role == "chapter" or role == "title" or "source-centered" in source_classes:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._append_inline(paragraph, text)
            paragraph.paragraph_format.page_break_before = page_break_before
            if role == "chapter-fragment":
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.space_after = Pt(5)
            return

        if node_type == "table":
            self._append_table(document, text, page_break_before=page_break_before)
            return

        if node_type in {"list", "list_item"}:
            self._append_list(document, node, text, page_break_before=page_break_before)
            return

        if node_type == "blockquote":
            for line_index, line in enumerate(text.splitlines() or [text]):
                content = re.sub(r"^\s*>\s?", "", line)
                paragraph = document.add_paragraph(style="Quote")
                if line_index == 0:
                    paragraph.paragraph_format.page_break_before = page_break_before
                self._append_inline(paragraph, content)
            return

        if node_type == "code":
            lines = text.splitlines()
            if lines and re.match(r"^\s*(?:```+|~~~+)", lines[0]):
                lines = lines[1:]
            if lines and re.match(r"^\s*(?:```+|~~~+)\s*$", lines[-1]):
                lines = lines[:-1]
            paragraph = document.add_paragraph(style="Code")
            paragraph.paragraph_format.page_break_before = page_break_before
            self._shade_paragraph(paragraph, "F2F2F2")
            for index, line in enumerate(lines):
                if index:
                    paragraph.add_run().add_break()
                run = paragraph.add_run(line)
                run.font.name = "Consolas"
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
            return

        self._append_text_node(
            document,
            node,
            role=role,
            metadata=metadata,
            source_classes=source_classes,
            page_break_before=page_break_before,
        )

    def _heading_style(self, metadata: dict[str, Any], role: str, source_classes: list[str]) -> str:
        if role == "title" or "source-title" in source_classes:
            return "Title"
        if role in {"contents", "contents-running"}:
            return "Contents Title"
        try:
            level = max(1, min(6, int(metadata.get("level") or 1)))
        except (TypeError, ValueError):
            level = 1
        if role == "chapter-fragment":
            level = max(2, level)
        return f"Heading {level}"

    def _append_text_node(
        self,
        document: Any,
        node: dict[str, Any],
        *,
        role: str,
        metadata: dict[str, Any],
        source_classes: list[str],
        page_break_before: bool,
    ) -> None:
        text = str(node.get("translated_text") or "")
        node_type = str(node.get("type") or "paragraph")
        is_footnote = "footnote" in source_classes
        is_contents_entry = role in {"contents", "contents-running"} or str(
            metadata.get("structure_type") or ""
        ).casefold() == "toc_entry"
        paragraphs = split_translated_paragraphs(node)
        if not paragraphs:
            paragraphs = [text]
        if len(paragraphs) == 1 and "\n\n" in paragraphs[0] and is_contents_entry:
            paragraphs = [line for line in re.split(r"(?:\r?\n[ \t]*){2,}", paragraphs[0]) if line.strip()]

        for paragraph_index, value in enumerate(paragraphs):
            lines = value.splitlines() or [""]
            if is_contents_entry:
                style = "Contents Entry 2" if source_toc_level(value) == 2 else "Contents Entry"
            elif is_footnote:
                style = "Footnote Text"
            elif role == "frontmatter":
                style = "Front Matter"
            else:
                style = "Normal"
            paragraph = document.add_paragraph(style=style)
            if paragraph_index == 0:
                paragraph.paragraph_format.page_break_before = page_break_before
            if "source-centered" in source_classes:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for line_index, line in enumerate(lines):
                if line_index:
                    paragraph.add_run().add_break()
                self._append_inline(paragraph, line)

            if metadata.get("page_number") and paragraph_index == 0:
                paragraph.paragraph_format.page_break_before = True

    def _append_list(
        self,
        document: Any,
        node: dict[str, Any],
        text: str,
        *,
        page_break_before: bool,
    ) -> None:
        metadata = node.get("attributes") if isinstance(node.get("attributes"), dict) else {}
        lines = text.splitlines() or [text]
        has_numeric_marker = any(
            match is not None and match.group("marker")[:1].isdigit()
            for match in (_LIST_MARKER.match(line) for line in lines)
        )
        ordered = bool(metadata.get("ordered")) if "ordered" in metadata else has_numeric_marker
        style = "List Number" if ordered else "List Bullet"
        next_number = 1
        shared_num_id = self._numbering_instance(document, next_number) if ordered and not has_numeric_marker else None
        for index, line in enumerate(lines):
            match = _LIST_MARKER.match(line)
            content = match.group("text") if match else line.strip()
            if not content:
                continue
            paragraph = document.add_paragraph(style=style)
            if index == 0:
                paragraph.paragraph_format.page_break_before = page_break_before
            if ordered:
                if match and match.group("marker")[:1].isdigit():
                    start_at = int(re.match(r"\d+", match.group("marker")).group(0))
                    num_id = self._numbering_instance(document, start_at)
                else:
                    num_id = shared_num_id or self._numbering_instance(document, next_number)
                    next_number += 1
                self._set_numbering(paragraph, num_id)
            self._append_inline(paragraph, content)

    @staticmethod
    def _numbering_instance(document: Any, start_at: int) -> int:
        numbering = document.part.numbering_part.element
        abstract_ids = [
            int(item.get(qn("w:abstractNumId")))
            for item in numbering.findall(qn("w:abstractNum"))
        ]
        num_ids = [
            int(item.get(qn("w:numId")))
            for item in numbering.findall(qn("w:num"))
        ]
        abstract_id = max(abstract_ids, default=-1) + 1
        num_id = max(num_ids, default=0) + 1

        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multilevel = OxmlElement("w:multiLevelType")
        multilevel.set(qn("w:val"), "singleLevel")
        abstract.append(multilevel)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        number_format = OxmlElement("w:numFmt")
        number_format.set(qn("w:val"), "decimal")
        level.append(number_format)
        label = OxmlElement("w:lvlText")
        label.set(qn("w:val"), "%1.")
        level.append(label)
        justification = OxmlElement("w:lvlJc")
        justification.set(qn("w:val"), "left")
        level.append(justification)
        abstract.append(level)
        first_num = next(iter(numbering.findall(qn("w:num"))), None)
        if first_num is None:
            numbering.append(abstract)
        else:
            numbering.insert(numbering.index(first_num), abstract)

        instance = OxmlElement("w:num")
        instance.set(qn("w:numId"), str(num_id))
        abstract_reference = OxmlElement("w:abstractNumId")
        abstract_reference.set(qn("w:val"), str(abstract_id))
        instance.append(abstract_reference)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        override_start = OxmlElement("w:startOverride")
        override_start.set(qn("w:val"), str(max(1, start_at)))
        override.append(override_start)
        instance.append(override)
        numbering.append(instance)
        return num_id

    @staticmethod
    def _set_numbering(paragraph: Any, num_id: int) -> None:
        ppr = paragraph._p.get_or_add_pPr()
        number_properties = OxmlElement("w:numPr")
        level = OxmlElement("w:ilvl")
        level.set(qn("w:val"), "0")
        number = OxmlElement("w:numId")
        number.set(qn("w:val"), str(num_id))
        number_properties.extend((level, number))
        ppr.append(number_properties)

    def _append_table(self, document: Any, text: str, *, page_break_before: bool) -> None:
        rows = self._table_rows(text)
        if not rows:
            if text.strip():
                paragraph = document.add_paragraph(style="Normal")
                paragraph.paragraph_format.page_break_before = page_break_before
                self._append_inline(paragraph, text)
            return
        column_count = max(len(row) for row in rows)
        table = document.add_table(rows=len(rows), cols=column_count)
        table.style = "Table Grid"
        table.autofit = True
        for row_index, row in enumerate(rows):
            for column_index in range(column_count):
                cell = table.cell(row_index, column_index)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                cell.text = ""
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                if row_index == 0:
                    self._shade_cell(cell, "EAEAEA")
                self._append_inline(paragraph, row[column_index] if column_index < len(row) else "", bold=row_index == 0)
        if page_break_before:
            # Tables do not expose paragraph_format; place the break on the
            # first cell paragraph while retaining the table as a Word table.
            table.cell(0, 0).paragraphs[0].paragraph_format.page_break_before = True

    @staticmethod
    def _table_rows(text: str) -> list[list[str]]:
        rows: list[list[str]] = []
        for line in text.splitlines():
            if line.strip() and "|" not in line:
                # A table node that mixes prose with pipe rows is ambiguous.
                # The caller keeps the complete translated block as text.
                return []
            if "|" not in line:
                continue
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if not cells or all(_TABLE_DIVIDER_CELL.fullmatch(cell) for cell in cells):
                continue
            rows.append(cells)
        return rows

    @classmethod
    def _append_inline(cls, paragraph: Any, value: str, *, bold: bool = False) -> None:
        cursor = 0
        for match in _INLINE_MARKUP.finditer(value):
            if match.start() > cursor:
                run = paragraph.add_run(value[cursor : match.start()])
                run.bold = bold
            kind = match.lastgroup
            raw = match.group(0)
            if kind == "strong":
                run = paragraph.add_run(raw[2:-2])
                run.bold = True
            elif kind == "emphasis":
                run = paragraph.add_run(raw[1:-1])
                run.italic = True
            else:
                run = paragraph.add_run(raw[1:-1])
                run.font.name = "Consolas"
            if bold:
                run.bold = True
            cursor = match.end()
        if cursor < len(value):
            run = paragraph.add_run(value[cursor:])
            run.bold = bold

    @staticmethod
    def _shade_paragraph(paragraph: Any, fill: str) -> None:
        ppr = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        ppr.append(shading)

    @staticmethod
    def _shade_cell(cell: Any, fill: str) -> None:
        tcpr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        tcpr.append(shading)

    @staticmethod
    def _source_style_classes(metadata: dict[str, Any]) -> list[str]:
        tokens = {
            str(value).strip().casefold()
            for value in metadata.get("class_tokens") or []
            if str(value).strip()
        }
        chapter_file = str(metadata.get("chapter_file") or "").replace("\\", "/")
        filename = chapter_file.rsplit("/", 1)[-1].casefold()
        structure_type = str(metadata.get("structure_type") or "").casefold()
        classes: list[str] = []
        if structure_type in {"footnote", "endnote"} or any(
            token == "footnote" or token.startswith(("footnote-", "endnote"))
            for token in tokens
        ):
            classes.append("footnote")
        if (
            any(token.startswith("center") for token in tokens)
            or tokens.intersection({"author", "dedication", "half-title"})
            or "copyright" in filename
            or any("copyright" in token for token in tokens)
        ):
            classes.append("source-centered")
        if "book-title" in tokens:
            classes.append("source-title")
        return classes

    @classmethod
    def _page_break_positions(
        cls,
        nodes: list[dict[str, Any]],
        roles: dict[str, str],
    ) -> set[int]:
        breaks: set[int] = set()
        copyright_positions = [
            index
            for index, node in enumerate(nodes)
            if cls._is_copyright(node)
        ]
        if copyright_positions:
            breaks.add(copyright_positions[0])
            after_copyright = copyright_positions[-1] + 1
            if after_copyright < len(nodes):
                breaks.add(after_copyright)

        previous_chapter_file = ""
        in_contents = False
        for index, node in enumerate(nodes):
            node_id = str(node.get("id") or "")
            role = roles.get(node_id)
            metadata = node.get("attributes") if isinstance(node.get("attributes"), dict) else {}
            chapter_file = str(metadata.get("chapter_file") or "")
            if role in {"contents", "contents-running"}:
                if not in_contents and index:
                    breaks.add(index)
                in_contents = True
            else:
                in_contents = False
            if node.get("type") == "heading" and role == "chapter":
                try:
                    level = int(metadata.get("level") or 1)
                except (TypeError, ValueError):
                    level = 1
                explicit_chapter = any(
                    str(metadata.get(key) or "").casefold() in {"chapter", "major"}
                    for key in ("epub_role", "layout_role", "display_role")
                ) or bool(metadata.get("chapter_break") or metadata.get("chapter_start"))
                changed_chapter_file = bool(chapter_file and previous_chapter_file and chapter_file != previous_chapter_file)
                if index and (level <= 1 or explicit_chapter or changed_chapter_file):
                    breaks.add(index)
            if metadata.get("page_number") and index:
                breaks.add(index)
            if chapter_file:
                previous_chapter_file = chapter_file
        return breaks

    @classmethod
    def _is_copyright(cls, node: dict[str, Any]) -> bool:
        metadata = node.get("attributes") if isinstance(node.get("attributes"), dict) else {}
        chapter_file = str(metadata.get("chapter_file") or "").replace("\\", "/").rsplit("/", 1)[-1].casefold()
        tokens = [str(value).casefold() for value in metadata.get("class_tokens") or []]
        return "copyright" in chapter_file or any("copyright" in token for token in tokens)

    def _target_path(self, state: dict[str, Any]) -> Path:
        target = (self.output_root / f"{self.assembler.source_stem(state)}_translated.docx").resolve()
        if target.parent != self.output_root:
            raise DocxExportError("DOCX 输出路径必须位于当前项目的 output 目录内。")
        return target

    @staticmethod
    def _validate_package(payload: bytes) -> None:
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                if archive.testzip() is not None:
                    raise DocxExportError("DOCX 压缩包完整性检查失败。")
        except zipfile.BadZipFile as exc:
            raise DocxExportError("DOCX 文件结构无效。") from exc
